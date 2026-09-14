from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
OUTPUT = ROOT / "coursework" / "courses" / "chem-58103" / "Syllabus" / "CHEM_481H3_CHEM_58103_Fall_2026_Syllabus.docx"


def set_cell_shading(cell, color):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), color)


def set_cell_margins(cell, top=95, start=110, bottom=95, end=110):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for margin, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{margin}"))
        if node is None:
            node = OxmlElement(f"w:{margin}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_cell_text(cell, text, bold=False, color="000000", size=9.5, align=WD_ALIGN_PARAGRAPH.LEFT):
    paragraph = cell.paragraphs[0]
    paragraph.alignment = align
    paragraph.paragraph_format.space_after = Pt(0)
    paragraph.paragraph_format.space_before = Pt(0)
    paragraph.paragraph_format.line_spacing = 1.08
    run = paragraph.add_run(text)
    run.bold = bold
    run.font.name = "Aptos"
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), "Aptos")
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), "Aptos")
    run.font.size = Pt(size)
    run.font.color.rgb = RGBColor.from_string(color)


def style_table(table, widths, header=True, exam_rows=None):
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    table.style = "Table Grid"
    exam_rows = exam_rows or set()
    for row_index, row in enumerate(table.rows):
        tr_pr = row._tr.get_or_add_trPr()
        cant_split = OxmlElement("w:cantSplit")
        tr_pr.append(cant_split)
        if row_index == 0 and header:
            tbl_header = OxmlElement("w:tblHeader")
            tbl_header.set(qn("w:val"), "true")
            tr_pr.append(tbl_header)
        for col_index, cell in enumerate(row.cells):
            cell.width = widths[col_index]
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            set_cell_margins(cell)
            if row_index == 0 and header:
                set_cell_shading(cell, "333F50")
            elif row_index in exam_rows:
                set_cell_shading(cell, "FCE8E6")
            elif row_index % 2 == 0:
                set_cell_shading(cell, "F3F6F9")
            for paragraph in cell.paragraphs:
                paragraph.paragraph_format.keep_together = True


def add_heading(doc, text, level=1):
    paragraph = doc.add_paragraph(style=f"Heading {level}")
    paragraph.paragraph_format.keep_with_next = True
    paragraph.paragraph_format.space_before = Pt(11 if level == 1 else 8)
    paragraph.paragraph_format.space_after = Pt(4)
    run = paragraph.add_run(text)
    run.font.color.rgb = RGBColor(0, 0, 0)
    return paragraph


def add_body(doc, text, bold_lead=None):
    paragraph = doc.add_paragraph()
    paragraph.paragraph_format.space_after = Pt(6)
    paragraph.paragraph_format.line_spacing = 1.12
    if bold_lead:
        lead = paragraph.add_run(bold_lead)
        lead.bold = True
    paragraph.add_run(text)
    return paragraph


doc = Document()
section = doc.sections[0]
section.page_width = Inches(8.5)
section.page_height = Inches(11)
section.top_margin = Inches(0.62)
section.bottom_margin = Inches(0.62)
section.left_margin = Inches(0.72)
section.right_margin = Inches(0.72)

styles = doc.styles
normal = styles["Normal"]
normal.font.name = "Aptos"
normal._element.rPr.rFonts.set(qn("w:ascii"), "Aptos")
normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Aptos")
normal.font.size = Pt(10.5)
normal.font.color.rgb = RGBColor(0, 0, 0)

title_style = styles["Title"]
title_style.font.name = "Aptos Display"
title_style._element.rPr.rFonts.set(qn("w:ascii"), "Aptos Display")
title_style._element.rPr.rFonts.set(qn("w:hAnsi"), "Aptos Display")
title_style.font.size = Pt(24)
title_style.font.bold = True
title_style.font.color.rgb = RGBColor(0, 0, 0)
# Some Word templates add a decorative lower border to the Title style. Keep the
# document clean and print-friendly by removing that inherited border explicitly.
title_style_ppr = title_style._element.get_or_add_pPr()
title_style_border = title_style_ppr.find(qn("w:pBdr"))
if title_style_border is not None:
    title_style_ppr.remove(title_style_border)

for style_name, size in (("Heading 1", 15), ("Heading 2", 12)):
    style = styles[style_name]
    style.font.name = "Aptos Display"
    style._element.rPr.rFonts.set(qn("w:ascii"), "Aptos Display")
    style._element.rPr.rFonts.set(qn("w:hAnsi"), "Aptos Display")
    style.font.size = Pt(size)
    style.font.bold = True
    style.font.color.rgb = RGBColor(0, 0, 0)

title = doc.add_paragraph(style="Title")
title.alignment = WD_ALIGN_PARAGRAPH.CENTER
title.paragraph_format.space_after = Pt(2)
title.add_run("CHEM 58103 Biochemistry I")

subtitle = doc.add_paragraph()
subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
subtitle.paragraph_format.space_after = Pt(11)
run = subtitle.add_run("Fall 2026 Syllabus")
run.bold = True
run.font.size = Pt(13)

info = doc.add_table(rows=5, cols=2)
info.style = "Table Grid"
info.alignment = WD_TABLE_ALIGNMENT.CENTER
info.autofit = False
labels = ["Course listing", "Instructor", "Email", "Class schedule", "Location"]
values = ["CHEM 58103 and CHEM 481H3", "Yizhi Sun, Assistant Professor of Biochemistry", "yizhis@uark.edu", "Monday, Wednesday, and Friday, 2:00-2:50 p.m.", "SCEN 402"]
for idx, (label, value) in enumerate(zip(labels, values)):
    set_cell_text(info.cell(idx, 0), label, bold=True, size=9.5)
    set_cell_text(info.cell(idx, 1), value, size=9.5)
style_table(info, [Inches(1.45), Inches(5.55)], header=False)
for row in info.rows:
    set_cell_shading(row.cells[0], "E8EEF4")

opening = doc.add_paragraph()
opening.paragraph_format.space_before = Pt(8)
opening.paragraph_format.space_after = Pt(7)
opening.paragraph_format.line_spacing = 1.12
opening.add_run("This syllabus uses CHEM 58103 as the primary course title. The instructor's paper syllabus lists the combined CHEM 481H3 and CHEM 58103 offering.")

add_heading(doc, "Course Description")
add_body(doc, "Biochemistry is the study of the chemistry of life processes. This course explores the basic molecules in biochemistry, including proteins, enzymes, DNA, RNA, carbohydrates, and lipids, as well as their association with human health.")

add_heading(doc, "Textbook")
add_body(doc, "Berg, Tymoczko, Gatto, and Stryer, Biochemistry, 9th edition.")

add_heading(doc, "Grading Plan")
grading = [
    ("Attendance", "10%"),
    ("Exam I", "20%"),
    ("Exam II", "20%"),
    ("Exam III", "20%"),
    ("Final Exam", "30%"),
]
table = doc.add_table(rows=1, cols=2)
set_cell_text(table.cell(0, 0), "Course component", bold=True, color="FFFFFF")
set_cell_text(table.cell(0, 1), "Weight", bold=True, color="FFFFFF", align=WD_ALIGN_PARAGRAPH.CENTER)
for component, weight in grading:
    cells = table.add_row().cells
    set_cell_text(cells[0], component)
    set_cell_text(cells[1], weight, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER)
style_table(table, [Inches(5.6), Inches(1.4)])

scale = [("A", "90-100%"), ("B", "75-89.9%"), ("C", "60-74.9%"), ("D", "40-59.9%"), ("F", "Below 40%")]
scale_table = doc.add_table(rows=2, cols=5)
for idx, (grade, value) in enumerate(scale):
    set_cell_text(scale_table.cell(0, idx), grade, bold=True, color="FFFFFF", align=WD_ALIGN_PARAGRAPH.CENTER)
    set_cell_text(scale_table.cell(1, idx), value, align=WD_ALIGN_PARAGRAPH.CENTER)
style_table(scale_table, [Inches(1.4)] * 5)

add_heading(doc, "Examination Policy")
add_body(doc, "Exams will be given during the regular class time. Makeup exams are available only to students who have a legitimate excuse for missing an exam, such as illness, a university-sanctioned event out of town, or a death in the immediate family. Keep documentation to support an anticipated absence and notify the instructor before the exam. If an exam is missed unexpectedly because of a last-minute illness or accident, retain the necessary documentation.")

add_heading(doc, "Academic Integrity")
add_body(doc, "As a core part of its mission, the University of Arkansas provides students with the opportunity to further their educational goals through programs of study and research in an environment that promotes freedom of inquiry and academic responsibility. Accomplishing this mission is only possible when intellectual honesty and individual integrity prevail. Each University of Arkansas student is required to be familiar with and abide by the University Academic Integrity Policy at http://honesty.uark.edu/policy.")

doc.add_page_break()

add_heading(doc, "Tentative and Actual Course Progress")
intro = doc.add_paragraph("The instructor may revise this schedule. The Actual Topics column will be updated from class records. Until a class meeting is documented, it remains marked Not yet recorded; no deviation is inferred.")
intro.paragraph_format.space_after = Pt(8)

schedule = [
    ("Week 1 (Aug. 17-21)", "1 and 2", "Introduction; amino acids and proteins", "Not yet recorded"),
    ("Week 2 (Aug. 24-28)", "3", "Protein analysis", "Not yet recorded"),
    ("Week 3 (Aug. 31-Sept. 4)", "4", "DNA and RNA", "Not yet recorded"),
    ("Week 4 (Sept. 7-11)", "5 and 6", "Gene and bioinformatics", "Not yet recorded"),
    ("Exam I (Sept. 11, 2:00-2:50 p.m.)", "2-6", "Exam I covering Chapters 2-6", "Assessment"),
    ("Week 5 (Sept. 14-18)", "7", "Hemoglobin", "Not yet recorded"),
    ("Week 6 (Sept. 21-25)", "8", "Enzymes: kinetics", "Not yet recorded"),
    ("Week 7 (Sept. 28-Oct. 2)", "9", "Enzymes: mechanism", "Not yet recorded"),
    ("Week 8 (Oct. 5-9)", "10", "Enzymes: regulation", "Not yet recorded"),
    ("Exam II (Oct. 9, 2:00-2:50 p.m.)", "7-10", "Exam II covering Chapters 7-10", "Assessment"),
    ("Week 9 (Oct. 12-16)", "11", "Carbohydrates", "Not yet recorded"),
    ("Week 10 (Oct. 19-23)", "12", "Lipid and membrane", "Not yet recorded"),
    ("Week 11 (Oct. 26-30)", "13", "Channels and pumps", "Not yet recorded"),
    ("Exam III (Oct. 30, 2:00-2:50 p.m.)", "11-13", "Exam III covering Chapters 11-13", "Assessment"),
    ("Week 12 (Nov. 2-6)", "28", "DNA replication", "Not yet recorded"),
    ("Week 13 (Nov. 9-13)", "29", "RNA synthesis and processing", "Not yet recorded"),
    ("Week 14 (Nov. 16-20)", "30", "Protein synthesis", "Not yet recorded"),
    ("Week 15 (Nov. 23-27)", "-", "No class", "-"),
    ("Week 16 (Nov. 30-Dec. 4)", "-", "Class review", "Not yet recorded"),
    ("Final Exam (Dec. 7-11, time TBD)", "All", "Final Exam covering all chapters", "Assessment"),
]

schedule_table = doc.add_table(rows=1, cols=4)
headers = ["Week (Dates)", "Chapters", "Tentative Topics", "Actual Topics"]
for idx, header in enumerate(headers):
    set_cell_text(schedule_table.cell(0, idx), header, bold=True, color="FFFFFF", size=9)
exam_rows = set()
for row_index, row_data in enumerate(schedule, start=1):
    cells = schedule_table.add_row().cells
    is_exam = row_data[0].startswith("Exam") or row_data[0].startswith("Final Exam")
    if is_exam:
        exam_rows.add(row_index)
    for col_index, value in enumerate(row_data):
        set_cell_text(cells[col_index], value, bold=is_exam, size=8.2, align=WD_ALIGN_PARAGRAPH.CENTER if col_index == 1 else WD_ALIGN_PARAGRAPH.LEFT)
style_table(schedule_table, [Inches(1.75), Inches(0.75), Inches(2.35), Inches(2.15)], exam_rows=exam_rows)

footer = section.footer.paragraphs[0]
footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
footer.paragraph_format.space_before = Pt(4)
footer_run = footer.add_run("CHEM 58103  |  Biochemistry I  |  Fall 2026  |  Also offered as CHEM 481H3")
footer_run.font.name = "Aptos"
footer_run.font.size = Pt(8)
footer_run.font.color.rgb = RGBColor(90, 90, 90)

doc.core_properties.title = "CHEM 58103 Biochemistry I Fall 2026 Syllabus"
doc.core_properties.subject = "Digitized course syllabus"
doc.core_properties.author = "Yizhi Sun"
doc.core_properties.keywords = "CHEM 481H3, CHEM 58103, Biochemistry I, Fall 2026"

OUTPUT.parent.mkdir(parents=True, exist_ok=True)
doc.save(OUTPUT)
print(OUTPUT)
